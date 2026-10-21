"""Acceptance test (SPEC §4): a whole day in mock mode, driven by a fake clock.

07:30 the morning brief arrives with drafts; 09:15 an urgent email interrupts; the user asks for work;
the task runs in a sandbox, asks for approval to email; the user edits and approves; the email goes
out; an artifact is produced; an unanswered draft expires and is denied.
"""

from __future__ import annotations

from datetime import timedelta
from pathlib import Path

from openpyxl import load_workbook

from tests.conftest import eventually, local


async def test_scripted_day(make_rt):
    rt = make_rt(start=local(7, 25))
    await rt.start()
    clock = rt.clock
    try:
        # 07:25 first heartbeat: check-in runs, the brief is not due yet.
        jobs = await rt.scheduler.tick()
        assert "checkin" in jobs and "brief" not in jobs
        assert rt.store.latest_brief() is None
        ignored = rt.store.list_events(status="ignored")
        assert any("cycling retail" in e["title"] for e in ignored), "the newsletter is ignored"

        # 07:30 the morning brief, with replies drafted as approval cards.
        clock.set(local(7, 30))
        jobs = await rt.scheduler.tick()
        assert "brief" in jobs
        brief = rt.store.latest_brief()
        assert brief and brief["greeting"].startswith("Good morning, Jordan")
        drafted = [i for i in brief["items"] if i.get("approval_id")]
        assert drafted, "at least one reply is drafted for review"
        draft_ap = rt.store.get_approval(drafted[0]["approval_id"])
        assert draft_ap["state"] == "pending" and draft_ap["tool"] == "email_send"
        cards = [m["card"]["type"] for m in rt.store.list_messages(limit=50) if m["card"]]
        assert "brief" in cards and "approval" in cards
        assert "brief" not in await rt.scheduler.tick(), "one brief per day"

        # 09:15 an urgent email arrives: a notification interrupts.
        clock.set(local(9, 15))
        pushed_before = len(rt.notifier.sent)
        await rt.scheduler.tick()
        pos = [e for e in rt.store.list_events(status="notified") if "POS subscription" in e["title"]]
        assert pos and pos[0]["urgency"] in ("high", "critical")
        assert len(rt.notifier.sent) > pushed_before
        assert any(n["title"].startswith("Payment failed") for n in rt.bus.of_type("notify"))

        # The user asks for real work in the thread.
        text = "Build me a one-page cash summary for this week and email it to Sam."
        rt.store.add_message("user", text)
        await rt.chat.reply(text)
        task = await eventually(lambda: next(iter(rt.store.list_tasks()), None))
        assert task["origin"] == "chat" and task["harness"] == "builtin"

        # The task works in its sandbox and stops at the approval gate before sending.
        ap = await eventually(
            lambda: next(
                (a for a in rt.store.list_approvals(state="pending") if a["task_id"] == task["id"]), None
            )
        )
        assert ap["tool"] == "email_send" and ap["args"]["to"] == "sam@ledgerwise.example"
        assert "body" in ap["editable"]
        assert rt.store.get_task(task["id"])["state"] == "waiting"
        assert not [o for o in rt.outbox() if o.get("to") == "sam@ledgerwise.example"], (
            "nothing sent before approval"
        )
        sandbox_files = list((Path(rt.settings.state_dir) / "sandboxes" / task["id"]).rglob("*"))
        assert any(p.name == "payments.csv" for p in sandbox_files)

        # The user edits the sign-off and approves: the action executes.
        edited_body = ap["args"]["body"] + "\nThanks, Jordan"
        await rt.approvals.decide(ap["id"], "edit", {"body": edited_body})
        done = await rt.tasks.wait(task["id"], timeout=20)
        assert done["state"] == "done", done
        sent = [o for o in rt.outbox() if o.get("to") == "sam@ledgerwise.example"]
        assert len(sent) == 1 and sent[0]["body"].endswith("Thanks, Jordan")
        audit = rt.store.list_audit(200)
        assert any(a["tool"] == "email_send" and a["decision"] == "edited" and a["ok"] == 1 for a in audit)
        assert any(m["source"] == "reflection" for m in rt.store.list_memory(["preference"])), (
            "learned from the edit"
        )

        # An artifact was produced and is a real workbook.
        assert done["artifact_ids"]
        art = rt.store.get_artifact(done["artifact_ids"][0])
        assert art["kind"] == "xlsx"
        wb = load_workbook(rt.artifacts.path(art))
        assert {"Summary", "Detail"} <= set(wb.sheetnames)
        assert "<table>" in rt.artifacts.preview(art)
        done_cards = [
            m for m in rt.store.list_messages(limit=100) if m["card"] and m["card"]["ref"] == art["id"]
        ]
        assert done_cards and done_cards[0]["card"]["type"] == "artifact"
        assert done_cards[0]["text"] == done["summary"]

        # Silence never approves: the brief's draft expires and is denied.
        clock.set(local(9, 15) + timedelta(seconds=rt.policy.timeout_s + 60))
        await rt.scheduler.tick()
        assert rt.store.get_approval(draft_ap["id"])["state"] == "expired"
        await eventually(
            lambda: any(
                a["approval_id"] == draft_ap["id"] and a["decision"] == "expired"
                for a in rt.store.list_audit(200)
            )
        )
        assert not [o for o in rt.outbox() if o.get("to") == draft_ap["args"]["to"]]

        # Today reflects the day.
        today = await rt.today()
        assert today["stats"]["tasks_done"] == 1
        assert today["brief"]["id"] == brief["id"]
        assert any(e["id"] == pos[0]["id"] for e in today["notifications"])
        assert today["upcoming"] is not None
    finally:
        await rt.stop()


async def test_triage_act_creates_task_and_quiet_hours_hold(make_rt):
    # 22:30: quiet hours. Run a check-in late at night: urgent-but-not-critical items wait for the brief.
    rt = make_rt(start=local(22, 30))
    await rt.start()
    try:
        await rt.checkin()
        overdue = next(e for e in rt.store.list_events() if "Overdue" in e["title"])
        assert overdue["status"] == "brief", "held during quiet hours"
        pos = next(e for e in rt.store.list_events() if "POS" in e["title"])
        assert pos["status"] == "notified" and pos["urgency"] == "critical", "critical always interrupts"
        priya = next(e for e in rt.store.list_events() if "lease" in e["title"].lower())
        assert priya["status"] == "acted"
        task = await eventually(
            lambda: next((t for t in rt.store.list_tasks() if t["origin"] == "triage"), None)
        )
        ap = await eventually(lambda: next(iter(rt.store.list_approvals(state="pending")), None))
        assert ap["args"]["to"] == "priya@riversideproperties.example"
        await rt.approvals.decide(ap["id"], "reject", note="I'll call her instead")
        done = await rt.tasks.wait(task["id"], timeout=20)
        assert done["state"] == "done" and "declined" in done["summary"]
        assert not rt.outbox()
    finally:
        await rt.stop()
