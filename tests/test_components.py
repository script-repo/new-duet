"""Connectors, sandbox, artifacts, memory and skills, triage rules, notification policy, scheduler."""

from __future__ import annotations

import json
import sys
from datetime import time, timedelta
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from tandem.artifacts import Artifacts
from tandem.config import (
    AppConfig,
    ConnectorConfig,
    ModelProfile,
    NotifyConfig,
    Settings,
    load_config,
    parse_duration,
)
from tandem.connectors import build_connectors
from tandem.connectors.calendar import parse_ics
from tandem.connectors.fake import FakeConnector
from tandem.connectors.mail import parse_message
from tandem.connectors.mcp import content_text, tags_for
from tandem.connectors.simple import FilesystemConnector, WebhookAuthError, WebhookConnector, sign
from tandem.engine.brief import brief_counts, brief_greeting
from tandem.engine.notify import NotificationPolicy, in_quiet_hours
from tandem.engine.scheduler import due
from tandem.engine.triage import apply_rules
from tandem.memory import Memory, Skills
from tandem.sandbox import LocalSandbox, SandboxError
from tandem.security import SecretResolver
from tandem.store import Store
from tandem.tools import ToolContext
from tandem.util import FakeClock, iso
from tests.conftest import TZ, local

# ---- fake connector ------------------------------------------------------------------------------------------


async def test_fake_connector_releases_events_over_the_day():
    clock = FakeClock(local(7, 0))
    fake = FakeConnector("demo", {"timezone": TZ}, clock)
    early = await fake.poll(None)
    titles = {e.title for e in early}
    assert "Q3 numbers for Thursday's board call?" in titles  # yesterday 22:10
    assert not any("POS" in t for t in titles)  # 09:10
    clock.set(local(9, 30))
    later = await fake.poll(None)
    assert any("POS" in e.title for e in later)
    assert len({e.external_id for e in later}) == len(later)


async def test_fake_connector_tools():
    clock = FakeClock(local(16, 0))
    fake = FakeConnector("demo", {"timezone": TZ}, clock)
    tools = {t.name: t for t in fake.tools()}
    assert {"email_send", "calendar_list", "payments_pay", "crm_lookup"} <= set(tools)
    assert "send" in tools["email_send"].tags and "pay" in tools["payments_pay"].tags
    ctx = ToolContext()
    found = await tools["email_search"].func({"query": "lease"}, ctx)
    assert found["messages"][0]["from"].startswith("priya")
    body = await tools["email_read"].func({"id": found["messages"][0]["id"]}, ctx)
    assert "lease renewal" in body["body"]
    pays = await tools["payments_list"].func({}, ctx)
    assert len(pays["payments"]) == 8
    for everything in ("all", "ALL", "any", "", None):
        assert len((await tools["payments_list"].func({"status": everything}, ctx))["payments"]) == 8
    overdue = await tools["payments_list"].func({"status": "overdue"}, ctx)
    assert [p["customer"] for p in overdue["payments"]] == ["Oakridge Cycling Club"]
    assert "all" in tools["payments_list"].parameters["properties"]["status"]["enum"]
    cal = await tools["calendar_list"].func({"days": 3}, ctx)
    assert cal["events"] and cal["free_slots"]
    assert (await tools["crm_lookup"].func({"name": "sam"}, ctx))["email"] == "sam@ledgerwise.example"
    await tools["email_send"].func({"to": "a@b.example", "subject": "s", "body": "b"}, ctx)
    await tools["payments_pay"].func({"payee": "X", "amount": 10}, ctx)
    assert [o["type"] for o in fake.outbox] == ["email", "payment"]
    items = await fake.upcoming(local(0, 0), local(23, 59))
    assert next(i.title for i in items) == "Team standup"


# ---- other connectors ---------------------------------------------------------------------------------------------

ICS = """BEGIN:VCALENDAR
BEGIN:VEVENT
UID:1
SUMMARY:Board call\\, Q3
DTSTART;TZID=America/Toronto:20261021T140000
DTEND;TZID=America/Toronto:20261021T150000
LOCATION:Zoom
ATTENDEE;CN=Mira:mailto:mira@brightwater.example
END:VEVENT
BEGIN:VEVENT
UID:2
SUMMARY:A very long title that is
  folded across lines
DTSTART:20261021T120000Z
END:VEVENT
BEGIN:VEVENT
UID:3
SUMMARY:Holiday
DTSTART;VALUE=DATE:20261022
END:VEVENT
END:VCALENDAR
"""


def test_parse_ics():
    items = parse_ics(ICS, TZ)
    assert [i.title for i in items] == [
        "A very long title that is folded across lines",
        "Board call, Q3",
        "Holiday",
    ]
    board = items[1]
    assert (
        board.start.hour == 14
        and board.location == "Zoom"
        and board.attendees == ["mira@brightwater.example"]
    )
    assert items[0].start.utcoffset() == timedelta(0)
    assert items[2].end - items[2].start == timedelta(days=1)


async def test_ics_connector_from_file(tmp_path):
    path = tmp_path / "cal.ics"
    path.write_text(ICS)
    from tandem.connectors.calendar import IcsConnector

    c = IcsConnector("calendar", {"path": str(path), "timezone": TZ}, FakeClock(local(13, 0)))
    up = await c.upcoming(local(0, 0), local(23, 0))
    assert any(i.title.startswith("Board call") for i in up)
    soon = await c.poll(None)
    assert any("Board call" in e.title for e in soon)
    assert [t.name for t in c.tools()] == ["calendar_list"]


async def test_webhook_hmac():
    secrets = SecretResolver({"HOOK": "s3cret-hook-value"})
    c = WebhookConnector("inbox", {"secret_env": "HOOK"}, FakeClock(local(9)), secrets)
    body = json.dumps({"title": "Deploy failed", "body": "prod is down", "id": "d1"}).encode()
    with pytest.raises(WebhookAuthError):
        await c.handle_webhook(json.loads(body), {}, body)
    events = await c.handle_webhook(
        json.loads(body), {"x-tandem-signature": sign("s3cret-hook-value", body)}, body
    )
    assert events[0].title == "Deploy failed" and events[0].external_id == "d1"
    open_hook = WebhookConnector("open", {}, FakeClock(local(9)))
    many = await open_hook.handle_webhook({"events": [{"subject": "a"}, {"text": "b"}]}, {})
    assert [e.title for e in many] == ["a", "b"]


async def test_filesystem_connector(tmp_path):
    root = tmp_path / "inbox"
    root.mkdir()
    (root / "old.txt").write_text("old")
    c = FilesystemConnector("files", {"root": str(root)}, FakeClock(local(9)))
    assert await c.poll(None) == []  # existing files are the baseline
    (root / "new.txt").write_text("new")
    events = await c.poll(None)
    assert [e.title for e in events] == ["New file: new.txt"]
    tools = {t.name: t for t in c.tools()}
    await tools["files_write"].func({"path": "sub/a.md", "content": "# hi"}, ToolContext())
    assert await tools["files_read"].func({"path": "sub/a.md"}, ToolContext()) == "# hi"
    with pytest.raises(PermissionError):
        await tools["files_read"].func({"path": "../secret"}, ToolContext())


def test_imap_message_parsing():
    raw = (
        b"From: Sam <sam@ledgerwise.example>\r\nSubject: Invoice\r\nDate: Wed, 21 Oct 2026 09:00:00 -0400\r\n"
        b"Message-ID: <abc@x>\r\nList-Unsubscribe: <mailto:u@x>\r\n\r\nPlease pay."
    )
    ev = parse_message(raw, "42", "email")
    assert ev.sender == "sam@ledgerwise.example" and ev.title == "Invoice" and ev.external_id == "<abc@x>"
    assert "Please pay." in ev.body and "List-Unsubscribe" in ev.data["headers"]


def test_mcp_tags_and_content():
    assert tags_for("search", {"readOnlyHint": True}, []) == frozenset({"read"})
    assert {"write", "delete"} <= tags_for("drop", {"destructiveHint": True}, [])
    assert "external_post" in tags_for("create_page", {}, ["*create*"])

    class C:
        text = "hello"

    class R:
        content = [C()]
        structuredContent = None

    assert content_text(R()) == "hello"


def test_build_connectors_validation():
    clock = FakeClock(local(9))
    with pytest.raises(ValueError, match="unknown connector"):
        build_connectors([ConnectorConfig(type="nope")], clock, None, TZ)
    with pytest.raises(ValueError, match="unique"):
        build_connectors(
            [ConnectorConfig(type="webhook", name="a"), ConnectorConfig(type="webhook", name="a")],
            clock,
            None,
            TZ,
        )
    assert build_connectors([ConnectorConfig(type="webhook", enabled=False)], clock, None, TZ) == {}


# ---- sandbox ---------------------------------------------------------------------------------------------------------


async def test_local_sandbox_confinement_and_exec(tmp_path, monkeypatch):
    monkeypatch.setenv("SOME_API_KEY", "should-not-leak-123")
    box = LocalSandbox(tmp_path / "box")
    box.write("a/b.txt", "hi")
    assert box.read("a/b.txt") == "hi" and box.list() == ["a/b.txt"]
    assert box.read("/work/a/b.txt") == "hi"
    for bad in ("../x", "../../etc/passwd", "a/../../x"):
        with pytest.raises(SandboxError):
            box.write(bad, "no")
    cmd = f'"{sys.executable}" -c "import os; print(os.getcwd()); print(os.environ.get(\'SOME_API_KEY\'))"'
    res = await box.exec(cmd)
    assert res.code == 0
    assert Path(res.stdout.splitlines()[0]).resolve() == box.root
    assert "should-not-leak" not in res.stdout


@pytest.mark.skipif(sys.platform == "win32", reason="sleep semantics differ")
async def test_local_sandbox_timeout(tmp_path):
    box = LocalSandbox(tmp_path / "box")
    res = await box.exec(f'"{sys.executable}" -c "import time; time.sleep(5)"', timeout=0.5)
    assert res.timed_out


# ---- artifacts ----------------------------------------------------------------------------------------------------------


def test_artifacts_every_kind(tmp_path):
    store = Store("sqlite://", clock=FakeClock(local(9)))
    arts = Artifacts(store, tmp_path / "artifacts")
    md = arts.create("md", "Notes", "# Hello\n\n- one\n- two")
    assert "<h1>Hello</h1>" in arts.preview(md)
    html = arts.create("html", "Page", "<p>hi <script>x</script></p>")
    assert "<p>hi" in arts.preview(html)
    docx = arts.create(
        "docx",
        "Memo",
        "# Memo\n\nIntro with **bold**.\n\n- point\n1. first\n\n| a | b |\n|---|---|\n| 1 | 2 |",
    )
    doc = Document(arts.path(docx))
    assert any("Intro with" in p.text for p in doc.paragraphs) and doc.tables
    assert "Intro with" in arts.preview(docx)
    xlsx = arts.create(
        "xlsx",
        "Sheet",
        json.dumps({"sheets": [{"name": "S", "rows": [["a", "b"], ["1", 2], ["", "=SUM(B2:B2)"]]}]}),
    )
    ws = load_workbook(arts.path(xlsx))["S"]
    assert ws["A2"].value == 1 and ws["B3"].value == "=SUM(B2:B2)"
    csvx = arts.create("xlsx", "From CSV", "x,y\n1,2")
    assert load_workbook(arts.path(csvx)).active["B2"].value == 2
    pptx = arts.create(
        "pptx",
        "Deck",
        json.dumps(
            {
                "slides": [
                    {"title": "T", "bullets": ["a"]},
                    {"title": "Two", "bullets": ["b", "c"], "notes": "n"},
                ]
            }
        ),
    )
    assert len(Presentation(arts.path(pptx)).slides) == 2
    assert "Slide 2" in arts.preview(pptx)
    md_deck = arts.create("pptx", "MD deck", "# One\n- a\n# Two\n- b")
    assert len(Presentation(arts.path(md_deck)).slides) == 2
    with pytest.raises(ValueError):
        arts.create("exe", "bad", "x")
    assert len(store.list_artifacts()) == 7


def test_register_file(tmp_path):
    store = Store("sqlite://", clock=FakeClock(local(9)))
    arts = Artifacts(store, tmp_path / "artifacts")
    src = tmp_path / "report.csv"
    src.write_text("a,b\n1,2")
    art = arts.register_file(src, "tk_1")
    assert art["kind"] == "csv" and "<table>" in arts.preview(art)


# ---- memory and skills ------------------------------------------------------------------------------------------------


def test_memory_tiers_rules_and_search():
    mem = Memory(Store("sqlite://", clock=FakeClock(local(9))), core_chars=200)
    mem.write("The user is Jordan.", "core")
    mem.write("vip: Mira@brightwater.example, sam@ledgerwise.example; mute: deals@x.example", "preference")
    mem.write("Supplier Trailhead raises prices Nov 1", "fact")
    mem.write("v1", "note", key="k")
    mem.write("v2", "note", key="k")
    assert "Jordan" in mem.core_block() and "Trailhead" not in mem.core_block()
    assert mem.rules("vip") == ["mira@brightwater.example", "sam@ledgerwise.example"]
    assert mem.rules("mute") == ["deals@x.example"]
    assert mem.search("trailhead prices")[0]["content"].startswith("Supplier")
    assert [m["content"] for m in mem.all() if m["key"] == "k"] == ["v2"]
    with pytest.raises(ValueError):
        mem.write("x", "bogus")


def test_skills(tmp_path):
    (tmp_path / "s1").mkdir()
    (tmp_path / "s1" / "SKILL.md").write_text("---\nname: one\ndescription: First skill\n---\nDo the thing.")
    (tmp_path / "two.md").write_text("Second skill, no front matter.\nMore.")
    skills = Skills([tmp_path, tmp_path / "missing"])
    assert {s["name"] for s in skills.public()} == {"one", "two"}
    assert skills.load("one") == "Do the thing."
    assert "two: Second skill" in skills.index()
    with pytest.raises(KeyError):
        skills.load("three")


# ---- triage, notifications, scheduler -------------------------------------------------------------------------------------


def test_triage_rules():
    vip, mute = ["mira@brightwater.example"], ["deals@adblast.example", "newsletter"]
    assert apply_rules({"title": "Hi", "sender": "mira@brightwater.example"}, vip, mute).floor == "high"
    assert (
        apply_rules({"title": "Sale", "sender": "deals@adblast.example"}, vip, mute).force_action == "ignore"
    )
    bulk = apply_rules({"title": "News", "data": {"headers": {"List-Unsubscribe": "<x>"}}}, vip, mute)
    assert "bulk" in bulk.flags and bulk.force_action == "ignore"
    money = apply_rules({"title": "Invoice overdue", "body": "$2,350.00 is past due"}, vip, mute)
    assert "money" in money.flags and money.floor == "high"
    sec = apply_rules({"title": "New sign-in to your account", "body": ""}, vip, mute)
    assert sec.floor == "high"
    assert apply_rules({"title": "Lunch?", "body": "tomorrow"}, vip, mute).floor == "low"


def _brief_event(rt, n: int, urgency: str = "normal") -> dict:
    ev = rt.store.add_event(
        {
            "source": "x",
            "kind": "email",
            "external_id": f"b{n}",
            "title": f"Item {n}",
            "sender": "a@x.example",
        }
    )
    return rt.store.update_event(ev["id"], status="brief", urgency=urgency, triage={"reason": f"r{n}"})


async def test_failed_brief_only_consumes_what_it_showed(make_rt):
    rt = make_rt()
    await rt.start()
    evs = [_brief_event(rt, n) for n in range(9)]
    brief_model = rt.models.for_role("brief")
    real = brief_model.complete_json

    async def broken(*a, **k):
        raise ValueError("used all max_tokens before answering")

    brief_model.complete_json = broken  # type: ignore[method-assign]
    first = await rt.brief.run()
    assert len(first["items"]) == 7
    status = {e["id"]: rt.store.get_event(e["id"])["status"] for e in evs}
    assert sorted(status.values()).count("briefed") == 7 and list(status.values()).count("brief") == 2
    left = {i for i, s in status.items() if s == "brief"}
    assert any("plain summary" in a["title"] for a in rt.store.list_activity())

    # The next brief still sees the two the fallback left out, even though they predate it.
    brief_model.complete_json = real  # type: ignore[method-assign]
    rt.clock.advance(hours=1)
    data = await rt.brief.gather(rt.clock.now())
    assert left <= {e["id"] for e in data["events"]}
    await rt.brief.run()
    assert all(rt.store.get_event(i)["status"] == "briefed" for i in left)
    await rt.stop()


async def test_brief_model_failure_before_delivery_consumes_nothing(make_rt):
    rt = make_rt()
    await rt.start()
    evs = [_brief_event(rt, n) for n in range(3)]

    def boom(*a, **k):
        raise RuntimeError("db down")

    rt.store.add_brief = boom  # type: ignore[method-assign]
    with pytest.raises(RuntimeError):
        await rt.brief.run()
    assert all(rt.store.get_event(e["id"])["status"] == "brief" for e in evs)
    await rt.stop()


def test_brief_greeting_counts_come_from_code():
    data = {"tasks_running": [{"id": "t1"}], "approvals": [{"id": "ap1"}, {"id": "ap2"}]}
    items = [
        {"topic": "a", "needs_you": True, "approval_id": "ap1"},
        {"topic": "b", "needs_you": True},
        {"topic": "c", "needs_you": False},
    ]
    g = brief_greeting("Good morning, Sam! 0 things need you today. Coffee first.", "Sam", data, items)
    assert g == "Good morning, Sam! Coffee first. 1 thing in progress and 3 that need you."
    assert brief_counts(data, items) == (1, 3)  # 2 flagged items + ap2, which no item carries
    assert brief_greeting("", None, {}, []) == "Good morning. 0 things in progress and 0 that need you."
    assert brief_greeting("Three new items and no rush", "Al", {}, [{"needs_you": True}]).endswith(
        "Good morning, Al. 0 things in progress and 1 that needs you."
    )


async def test_brief_greeting_matches_its_items_and_counts_triage_tasks(make_rt):
    rt = make_rt()
    await rt.start()
    _brief_event(rt, 1, "high")
    _brief_event(rt, 2)
    task = rt.store.create_task("Reschedule the VIP call", "x", "builtin", origin="triage")
    rt.store.update_task(task["id"], state="running")
    acted = rt.store.add_event(
        {
            "source": "x",
            "kind": "email",
            "external_id": "vip",
            "title": "From the VIP",
            "sender": "vip@x.example",
        }
    )
    rt.store.update_event(acted["id"], status="acted", urgency="high", triage={"reason": "handled"})
    data = await rt.brief.gather(rt.clock.now())
    assert acted["id"] in {e["id"] for e in data["events"]}

    async def lying(*a, **k):
        return {
            "greeting": "Morning! 0 things in progress and 0 that need you.",
            "items": [
                {"topic": "Late fee", "summary": "Pay today.", "needs_you": True},
                {"topic": "VIP", "summary": "On it.", "needs_you": True},
                {"topic": "FYI", "summary": "Nothing to do.", "needs_you": False},
            ],
        }

    rt.models.for_role("brief").complete_json = lying  # type: ignore[method-assign]
    brief = await rt.brief.run()
    pending = len(rt.store.list_approvals(state="pending"))
    assert brief["greeting"] == f"Morning! 1 thing in progress and {2 + pending} that need you."
    await rt.stop()


async def test_restart_requeues_queued_tasks_and_fails_interrupted_ones(make_rt):
    rt = make_rt()
    queued = rt.store.create_task("Write a memo", "Write a short memo about the week", "builtin")
    running = rt.store.create_task("Half done", "x", "builtin")
    rt.store.update_task(running["id"], state="running")
    waiting = rt.store.create_task("Waiting on OK", "x", "builtin")
    rt.store.update_task(waiting["id"], state="waiting")
    gone = rt.store.create_task("Old harness", "x", "nope")
    await rt.start()
    for t in (running, waiting, gone):
        got = rt.store.get_task(t["id"])
        assert got["state"] == "failed" and got["error"].startswith("server restarted"), got
        assert got["finished_at"]
    done = await rt.tasks.wait(queued["id"], 20)
    assert done["state"] not in ("queued", "failed"), done
    assert any("requeued" in a["title"] for a in rt.store.list_activity())
    await rt.stop()


async def test_triage_falls_back_to_rules_when_model_breaks(make_rt):
    rt = make_rt()
    await rt.start()

    async def broken(*a, **k):
        raise RuntimeError("model down")

    rt.models.for_role("triage").complete_json = broken  # type: ignore[method-assign]
    ev = rt.store.add_event(
        {
            "source": "x",
            "kind": "email",
            "external_id": "1",
            "title": "Card declined",
            "body": "$50.00 payment failed, due today",
            "sender": "billing@x.example",
        }
    )
    verdicts = await rt.triage.run([ev])
    assert verdicts[0].source == "rules" and verdicts[0].urgency == "high" and verdicts[0].action == "notify"
    assert any("rules only" in a["title"] for a in rt.store.list_activity())
    await rt.stop()


def test_quiet_hours_and_policy():
    assert in_quiet_hours(local(22, 0), "21:00-07:00") and in_quiet_hours(local(6, 59), "21:00-07:00")
    assert not in_quiet_hours(local(7, 0), "21:00-07:00")
    assert in_quiet_hours(local(13, 0), "12:00-14:00") and not in_quiet_hours(local(14, 0), "12:00-14:00")
    store = Store("sqlite://", clock=FakeClock(local(9)))
    pol = NotificationPolicy(NotifyConfig(max_per_hour=2), TZ, store)
    assert pol.decide("normal", local(10)) == "hold"
    assert pol.decide("high", local(23)) == "hold"
    assert pol.decide("critical", local(23)) == "push"
    assert pol.decide("high", local(10)) == "push"
    assert pol.decide("high", local(10, 5)) == "push"
    assert pol.decide("high", local(10, 9)) == "coalesce"  # third in the same hour


def test_scheduler_due():
    t = time(7, 30)
    every = timedelta(minutes=15)
    assert set(due(local(7, 0), {}, every, t, TZ)) == {"checkin", "expire_approvals"}
    state = {
        "last_checkin": iso(local(7, 0)),
        "last_expiry": iso(
            local(
                7,
                29,
            )
        ),
    }
    assert due(local(7, 10), state, every, t, TZ) == []
    assert due(local(7, 30), state, every, t, TZ) == ["checkin", "brief", "expire_approvals"]
    state["last_brief_date"] = "2026-10-21"
    assert "brief" not in due(local(8, 0), state, every, t, TZ)
    assert "brief" in due(local(7, 31, day=22), state, every, t, TZ)


def test_config_loading(tmp_path, monkeypatch):
    for k in ("ANTHROPIC_API_KEY", "OPENAI_API_KEY", "GEMINI_API_KEY", "TANDEM_MOCK"):
        monkeypatch.delenv(k, raising=False)
    cfg = load_config(Settings(data_dir=tmp_path))
    assert cfg.mock and [c.type for c in cfg.connectors] == ["fake", "webhook"] and cfg.demo_autoplay
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test")
    cfg = load_config(Settings(data_dir=tmp_path))
    assert not cfg.mock and cfg.profile_for("triage").model == "gpt-5-mini"
    path = tmp_path / "t.yaml"
    path.write_text(
        "models:\n  big: {provider: gemini, model: g, api_key_env: G}\nroles: {worker: big}\nschedule: {brief_at: '06:45'}\n"
    )
    cfg = load_config(Settings(data_dir=tmp_path, config=path, timezone="Europe/Paris"))
    assert cfg.profile_for("worker").provider == "gemini" and cfg.schedule.brief_time == time(6, 45)
    assert cfg.schedule.timezone == "Europe/Paris"
    assert load_config(Settings(data_dir=tmp_path, config=path, mock=True)).mock
    assert parse_duration("1h30m") == timedelta(minutes=90) and parse_duration(45) == timedelta(seconds=45)
    assert AppConfig(models={"x": ModelProfile(provider="mock")}).profile_for("brief").name == "x"
