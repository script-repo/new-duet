"""Deliverables (SPEC §3.12): markdown, HTML, Word, Excel and PowerPoint files, plus HTML previews."""

from __future__ import annotations

import csv
import html
import io
import json
import re
from pathlib import Path
from typing import Any

from markdown_it import MarkdownIt

from tandem.store import Store
from tandem.util import new_id

KINDS = ("md", "html", "docx", "xlsx", "pptx")
ARTIFACT_EXTS = {
    ".md": "md",
    ".html": "html",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".pptx": "pptx",
    ".csv": "csv",
    ".pdf": "pdf",
    ".png": "png",
    ".svg": "svg",
}
_md = MarkdownIt("commonmark", {"html": False}).enable("table")

PAGE_CSS = """
body{font:16px/1.55 system-ui,-apple-system,Segoe UI,Roboto,sans-serif;max-width:860px;margin:2rem auto;padding:0 1rem;color:#1f1d1a;background:#fffdf9}
h1,h2,h3{line-height:1.25} table{border-collapse:collapse;margin:1rem 0} td,th{border:1px solid #ddd6cc;padding:.35rem .6rem;text-align:left}
th{background:#f4efe7} code{background:#f4efe7;padding:.1rem .3rem;border-radius:4px}
.slide{border:1px solid #ddd6cc;border-radius:12px;padding:1rem 1.25rem;margin:1rem 0;background:#fff}
.slide h2{margin-top:0} .muted{color:#6b645b;font-size:.9em}
@media (prefers-color-scheme: dark){body{background:#1b1a18;color:#ece7df} td,th{border-color:#3a3631} th,code{background:#2a2724} .slide{background:#23211e;border-color:#3a3631}}
"""


def slug(text: str) -> str:
    s = re.sub(r"[^A-Za-z0-9]+", "-", text).strip("-").lower()
    return (s or "artifact")[:60]


def page(title: str, body_html: str) -> str:
    return (
        f"<!doctype html><html><head><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'>"
        f"<title>{html.escape(title)}</title><style>{PAGE_CSS}</style></head><body>{body_html}</body></html>"
    )


def md_to_html(text: str) -> str:
    return _md.render(text)


# ---- writers -------------------------------------------------------------------------------------------------


def _inline_runs(paragraph: Any, text: str) -> None:
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if chunk:
            run = paragraph.add_run(chunk)
            run.bold = i % 2 == 1


def write_docx(path: Path, title: str, content: str) -> None:
    from docx import Document

    doc = Document()
    doc.core_properties.title = title
    lines = content.splitlines()
    if not lines or not lines[0].startswith("# "):
        doc.add_heading(title, 0)
    table_rows: list[list[str]] = []

    def flush_table() -> None:
        if not table_rows:
            return
        rows = [r for r in table_rows if not all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in r)]
        if rows:
            t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            t.style = "Table Grid"
            for i, r in enumerate(rows):
                for j, c in enumerate(r):
                    t.cell(i, j).text = c.strip()
        table_rows.clear()

    for line in lines:
        s = line.rstrip()
        if s.startswith("|") and s.endswith("|"):
            table_rows.append(s.strip("|").split("|"))
            continue
        flush_table()
        if not s.strip():
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", s)
        if m:
            doc.add_heading(m.group(2), 0 if len(m.group(1)) == 1 else len(m.group(1)) - 1)
        elif re.match(r"^\s*[-*]\s+", s):
            _inline_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", s))
        elif re.match(r"^\s*\d+[.)]\s+", s):
            _inline_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", s))
        else:
            _inline_runs(doc.add_paragraph(), s)
    flush_table()
    doc.save(str(path))


def _sheets_from(content: Any) -> list[dict[str, Any]]:
    if isinstance(content, str):
        try:
            content = json.loads(content)
        except json.JSONDecodeError:
            rows = list(csv.reader(io.StringIO(content)))
            return [{"name": "Sheet1", "rows": rows}]
    if isinstance(content, dict) and "sheets" in content:
        return content["sheets"]
    if isinstance(content, dict) and "rows" in content:
        return [{"name": content.get("name", "Sheet1"), "rows": content["rows"]}]
    if isinstance(content, list):
        return [{"name": "Sheet1", "rows": content}]
    raise ValueError("xlsx content must be {'sheets': [{'name', 'rows'}]}, rows, or CSV text")


def _cell(value: Any) -> Any:
    if isinstance(value, str):
        v = value.strip()
        if re.fullmatch(r"-?\d+(\.\d+)?", v):
            return float(v) if "." in v else int(v)
    return value


def write_xlsx(path: Path, title: str, content: Any) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    wb.remove(wb.active)
    for sheet in _sheets_from(content):
        ws = wb.create_sheet(str(sheet.get("name", "Sheet"))[:31])
        for r, row in enumerate(sheet.get("rows", []), start=1):
            for c, val in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c, value=_cell(val))
                if r == 1:
                    cell.font = Font(bold=True)
                    cell.fill = PatternFill("solid", fgColor="F4EFE7")
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 60)
        ws.freeze_panes = "A2"
    wb.properties.title = title
    wb.save(str(path))


def _slides_from(content: Any, title: str) -> list[dict[str, Any]]:
    if isinstance(content, str):
        try:
            content = json.loads(content)
        except json.JSONDecodeError:
            slides: list[dict[str, Any]] = []
            for line in content.splitlines():
                if line.startswith("#"):
                    slides.append({"title": line.lstrip("#").strip(), "bullets": []})
                elif line.strip() and slides:
                    slides[-1]["bullets"].append(re.sub(r"^\s*[-*\d.)]+\s*", "", line).strip())
            return slides or [{"title": title, "bullets": [content[:200]]}]
    if isinstance(content, dict) and "slides" in content:
        return content["slides"]
    if isinstance(content, list):
        return content
    raise ValueError("pptx content must be {'slides': [{'title', 'bullets'}]} or markdown with # per slide")


def write_pptx(path: Path, title: str, content: Any) -> None:
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    for i, s in enumerate(_slides_from(content, title)):
        bullets = [str(b) for b in s.get("bullets", [])]
        if i == 0:
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = str(s.get("title", title))
            if len(slide.placeholders) > 1:
                slide.placeholders[1].text = "\n".join(bullets)
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = str(s.get("title", ""))
            body = slide.placeholders[1].text_frame
            body.clear()
            for j, b in enumerate(bullets):
                p = body.paragraphs[0] if j == 0 else body.add_paragraph()
                p.text = b
                p.font.size = Pt(22)
        if s.get("notes"):
            slide.notes_slide.notes_text_frame.text = str(s["notes"])
    prs.core_properties.title = title
    prs.save(str(path))


# ---- previews --------------------------------------------------------------------------------------------------


def preview_html(path: Path, kind: str, title: str) -> str:
    try:
        if kind == "md":
            return page(title, md_to_html(path.read_text(encoding="utf-8")))
        if kind == "html":
            return path.read_text(encoding="utf-8")
        if kind == "csv":
            rows = list(csv.reader(io.StringIO(path.read_text(encoding="utf-8"))))
            return page(title, f"<h1>{html.escape(title)}</h1>" + _table(rows[:200]))
        if kind == "docx":
            from docx import Document

            doc = Document(str(path))
            parts = [f"<h1>{html.escape(title)}</h1>"]
            for p in doc.paragraphs:
                if not p.text.strip():
                    continue
                style = (p.style.name or "").lower() if p.style is not None else ""
                tag = "h2" if style.startswith(("heading", "title")) else "li" if "list" in style else "p"
                parts.append(f"<{tag}>{html.escape(p.text)}</{tag}>")
            for t in doc.tables:
                parts.append(_table([[c.text for c in r.cells] for r in t.rows]))
            return page(title, "\n".join(parts))
        if kind == "xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(str(path), read_only=True)
            parts = [f"<h1>{html.escape(title)}</h1>"]
            for ws in wb.worksheets:
                rows = [
                    ["" if v is None else v for v in r] for r in ws.iter_rows(max_row=200, values_only=True)
                ]
                parts.append(f"<h2>{html.escape(ws.title)}</h2>" + _table(rows))
            return page(title, "\n".join(parts))
        if kind == "pptx":
            from pptx import Presentation

            prs = Presentation(str(path))
            parts = [f"<h1>{html.escape(title)}</h1>"]
            for i, slide in enumerate(prs.slides, start=1):
                texts = [
                    sh.text_frame.text
                    for sh in slide.shapes
                    if sh.has_text_frame and sh.text_frame.text.strip()
                ]
                head = html.escape(texts[0]) if texts else f"Slide {i}"
                items = "".join(
                    f"<li>{html.escape(line)}</li>"
                    for t in texts[1:]
                    for line in t.splitlines()
                    if line.strip()
                )
                parts.append(
                    f"<div class='slide'><div class='muted'>Slide {i}</div><h2>{head}</h2><ul>{items}</ul></div>"
                )
            return page(title, "\n".join(parts))
    except Exception as exc:  # a broken file should not break the page
        return page(title, f"<p>Preview unavailable: {html.escape(str(exc))}</p>")
    return page(title, f"<p>No preview for .{html.escape(kind)} files. Download it instead.</p>")


def _table(rows: list[list[Any]]) -> str:
    if not rows:
        return "<p class='muted'>(empty)</p>"
    head = "".join(f"<th>{html.escape(str(c))}</th>" for c in rows[0])
    body = "".join("<tr>" + "".join(f"<td>{html.escape(str(c))}</td>" for c in r) + "</tr>" for r in rows[1:])
    return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"


class Artifacts:
    def __init__(self, store: Store, base: Path) -> None:
        self.store = store
        self.base = base

    def create(self, kind: str, title: str, content: Any, task_id: str | None = None) -> dict[str, Any]:
        kind = kind.lower().lstrip(".")
        if kind not in KINDS:
            raise ValueError(f"kind must be one of {', '.join(KINDS)}")
        aid = new_id("ar")
        folder = self.base / aid
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{slug(title)}.{kind}"
        text = content if isinstance(content, str) else json.dumps(content)
        if kind == "md":
            path.write_text(text, encoding="utf-8")
        elif kind == "html":
            body = text if re.search(r"<\w+", text) else md_to_html(text)
            path.write_text(body if "<html" in body.lower() else page(title, body), encoding="utf-8")
        elif kind == "docx":
            write_docx(path, title, text)
        elif kind == "xlsx":
            write_xlsx(path, title, content)
        elif kind == "pptx":
            write_pptx(path, title, content)
        return self.store.add_artifact(
            id=aid,
            task_id=task_id,
            kind=kind,
            title=title,
            filename=path.name,
            path=str(path),
            size=path.stat().st_size,
        )

    def register_file(self, src: Path, task_id: str | None, title: str | None = None) -> dict[str, Any]:
        """Copy a file a harness produced (e.g. a CLI agent's output) into the library."""
        kind = ARTIFACT_EXTS.get(src.suffix.lower(), src.suffix.lower().lstrip(".") or "bin")
        aid = new_id("ar")
        folder = self.base / aid
        folder.mkdir(parents=True, exist_ok=True)
        dest = folder / src.name
        dest.write_bytes(src.read_bytes())
        return self.store.add_artifact(
            id=aid,
            task_id=task_id,
            kind=kind,
            title=title or src.stem.replace("-", " "),
            filename=src.name,
            path=str(dest),
            size=dest.stat().st_size,
        )

    def path(self, artifact: dict[str, Any]) -> Path:
        p = Path(artifact["path"]).resolve()
        if self.base.resolve() not in p.parents:
            raise ValueError("artifact path outside the library")
        return p

    def preview(self, artifact: dict[str, Any]) -> str:
        return preview_html(self.path(artifact), artifact["kind"], artifact["title"])
